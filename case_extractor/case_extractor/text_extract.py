"""PDF / DOCX / DOC 판결문 파일에서 텍스트를 추출한다."""

from __future__ import annotations

import re
import struct
from pathlib import Path


class TextExtractError(Exception):
    pass


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".doc":
        return _extract_doc(path)
    raise TextExtractError(
        f"{path.name}: 지원하지 않는 파일 형식입니다 (.pdf, .docx, .doc만 지원)."
    )


def _extract_pdf(path: Path) -> str:
    import pdfplumber

    chunks: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            chunks.append(text)
    full_text = "\n".join(chunks).strip()
    if not full_text:
        raise TextExtractError(
            f"{path.name}: 텍스트를 추출하지 못했습니다. 스캔된 이미지 PDF일 수 있습니다 "
            "(OCR이 필요합니다. macOS에서는 'ocrmypdf'로 먼저 텍스트 레이어를 입힌 뒤 다시 시도하세요)."
        )
    return full_text


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    full_text = "\n".join(parts).strip()
    if not full_text:
        raise TextExtractError(f"{path.name}: 문서에서 텍스트를 찾을 수 없습니다.")
    return full_text


# ── .doc (Word 97-2003 바이너리) ──────────────────────────────────────
# OLE 복합문서 안의 WordDocument 스트림을 piece table(CLX)로 해석해 텍스트를 뽑는다.

_FIB_IDENT = 0xA5EC
_OFF_FLAGS = 0x000A      # fWhichTblStm = bit 9
_OFF_FC_CLX = 0x01A2
_OFF_LCB_CLX = 0x01A6
_CLXT_PRC = 0x01
_CLXT_PCDT = 0x02
_FC_COMPRESSED = 0x40000000
_FC_MASK = 0x3FFFFFFF


def _extract_doc(path: Path) -> str:
    try:
        import olefile
    except ImportError as e:
        raise TextExtractError(
            "olefile 패키지가 필요합니다. 'pip install -r requirements.txt'를 실행하세요."
        ) from e

    if not olefile.isOleFile(str(path)):
        # 확장자만 .doc 이고 실제로는 .docx(zip) 인 경우가 흔하다
        try:
            return _extract_docx(path)
        except Exception:
            pass
        raise TextExtractError(
            f"{path.name}: 올바른 Word 97-2003(.doc) 파일이 아닙니다. "
            "한글/워드에서 '.docx'로 다시 저장한 뒤 시도하세요."
        )

    with olefile.OleFileIO(str(path)) as ole:
        if not ole.exists("WordDocument"):
            raise TextExtractError(
                f"{path.name}: Word 문서 구조를 찾을 수 없습니다 "
                "(Excel/PowerPoint 파일이거나 손상된 파일일 수 있습니다)."
            )
        doc_stream = ole.openstream("WordDocument").read()

        if len(doc_stream) < _OFF_LCB_CLX + 4:
            raise TextExtractError(f"{path.name}: 파일이 손상되었거나 내용이 없습니다.")
        if struct.unpack_from("<H", doc_stream, 0)[0] != _FIB_IDENT:
            raise TextExtractError(f"{path.name}: Word 문서 서명이 올바르지 않습니다.")

        flags = struct.unpack_from("<H", doc_stream, _OFF_FLAGS)[0]
        table_name = "1Table" if (flags >> 9) & 1 else "0Table"
        if not ole.exists(table_name):
            table_name = "0Table" if table_name == "1Table" else "1Table"
        if not ole.exists(table_name):
            raise TextExtractError(f"{path.name}: 문서 테이블 스트림을 찾을 수 없습니다.")
        table_stream = ole.openstream(table_name).read()

    fc_clx = struct.unpack_from("<I", doc_stream, _OFF_FC_CLX)[0]
    lcb_clx = struct.unpack_from("<I", doc_stream, _OFF_LCB_CLX)[0]
    if lcb_clx == 0 or fc_clx + lcb_clx > len(table_stream):
        raise TextExtractError(
            f"{path.name}: 텍스트 위치 정보를 읽지 못했습니다. "
            "한글/워드에서 '.docx'로 다시 저장한 뒤 시도하세요."
        )

    pieces = _parse_piece_table(table_stream[fc_clx : fc_clx + lcb_clx])
    if not pieces:
        raise TextExtractError(f"{path.name}: 문서에서 텍스트 조각을 찾지 못했습니다.")

    parts: list[str] = []
    for cp_start, cp_end, fc in pieces:
        n = cp_end - cp_start
        if n <= 0:
            continue
        if fc & _FC_COMPRESSED:
            offset = (fc & _FC_MASK) // 2
            raw = doc_stream[offset : offset + n]
            parts.append(raw.decode("cp1252", errors="replace"))
        else:
            offset = fc & _FC_MASK
            raw = doc_stream[offset : offset + n * 2]
            parts.append(raw.decode("utf-16-le", errors="replace"))

    text = _clean_doc_text("".join(parts))
    if not text:
        raise TextExtractError(
            f"{path.name}: 텍스트를 추출하지 못했습니다. "
            "한글/워드에서 '.docx'로 다시 저장한 뒤 시도하세요."
        )
    return text


def _parse_piece_table(clx: bytes) -> list[tuple[int, int, int]]:
    """CLX에서 Pcdt를 찾아 [(cpStart, cpEnd, fc), ...] 목록을 만든다."""
    i = 0
    while i < len(clx):
        kind = clx[i]
        if kind == _CLXT_PRC:
            # Prc: clxt(1) + cbGrpprl(2) + grpprl(cbGrpprl) — 건너뛴다
            if i + 3 > len(clx):
                break
            cb = struct.unpack_from("<h", clx, i + 1)[0]
            i += 3 + max(cb, 0)
        elif kind == _CLXT_PCDT:
            if i + 5 > len(clx):
                break
            lcb = struct.unpack_from("<I", clx, i + 1)[0]
            plc = clx[i + 5 : i + 5 + lcb]
            return _parse_plcpcd(plc)
        else:
            break
    return []


def _parse_plcpcd(plc: bytes) -> list[tuple[int, int, int]]:
    """PlcPcd = CP 배열(n+1개, 4바이트) + PCD 배열(n개, 8바이트)."""
    if len(plc) < 4 + 8:
        return []
    n = (len(plc) - 4) // 12
    if n <= 0:
        return []

    cps = [struct.unpack_from("<I", plc, k * 4)[0] for k in range(n + 1)]
    pcd_base = (n + 1) * 4

    pieces: list[tuple[int, int, int]] = []
    for k in range(n):
        fc = struct.unpack_from("<I", plc, pcd_base + k * 8 + 2)[0]
        pieces.append((cps[k], cps[k + 1], fc))
    return pieces


# 필드 코드: \x13 지시문 \x14 결과 \x15 → 결과만 남긴다
_FIELD_RE = re.compile(r"\x13[^\x14\x15]*(?:\x14([^\x15]*))?\x15")
# 남는 제어문자 (문단/줄바꿈/탭은 따로 처리)
_CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0e-\x1f]")


def _clean_doc_text(raw: str) -> str:
    text = _FIELD_RE.sub(lambda m: m.group(1) or "", raw)
    text = text.replace("\x07", "\n")   # 표 셀/행 구분
    text = text.replace("\r", "\n")     # 문단 끝
    text = text.replace("\x0c", "\n")   # 페이지/구역 구분
    text = text.replace("\xa0", " ")    # 줄바꿈 없는 공백
    text = _CTRL_RE.sub("", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".doc")


def find_case_files(input_dir: Path) -> list[Path]:
    files = [
        p for p in sorted(input_dir.iterdir())
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    ]
    return files
