"""추출된 행을 사용자의 코딩시트 엑셀 템플릿에 이어서 기록한다."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .coding_book import SHEET_COLUMN_ORDER

SHEET_NAME = "코딩시트"
NOTES_SHEET_NAME = "코딩노트"
EXAMPLE_ROW_MARKER = "예시행"


def _next_case_number(ws, id_prefix: str) -> int:
    pattern = re.compile(re.escape(id_prefix) + r"(\d+)$")
    max_n = 0
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        value = row[0]
        if not value:
            continue
        m = pattern.match(str(value).strip())
        if m:
            max_n = max(max_n, int(m.group(1)))
    return max_n + 1


def _find_insert_row(ws) -> int:
    """헤더/예시행 다음, 데이터가 없는 첫 행 번호를 찾는다."""
    last = 1
    for row_idx in range(2, ws.max_row + 1):
        case_id = ws.cell(row=row_idx, column=1).value
        note = ws.cell(row=row_idx, column=len(SHEET_COLUMN_ORDER) + 1).value
        is_example = bool(note) and EXAMPLE_ROW_MARKER in str(note)
        if case_id and not is_example:
            last = row_idx
        elif is_example:
            continue
    return last + 1


def write_rows(
    template_path: Path,
    output_path: Path,
    rows: list[dict],
    *,
    id_prefix: str = "DF-2024-",
) -> Path:
    """rows(각 항목이 SHEET_COLUMN_ORDER 키를 갖는 dict)를 템플릿에 이어붙여 output_path에 저장한다."""
    if output_path != template_path:
        shutil.copy(template_path, output_path)
        target = output_path
    else:
        target = template_path

    wb = openpyxl.load_workbook(target)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"'{SHEET_NAME}' 시트를 템플릿 파일에서 찾을 수 없습니다: {template_path}")
    ws = wb[SHEET_NAME]

    next_row = _find_insert_row(ws)
    next_num = _next_case_number(ws, id_prefix)

    for row in rows:
        if not row.get("case_id"):
            row["case_id"] = f"{id_prefix}{next_num:03d}"
            next_num += 1
        for col_idx, col_name in enumerate(SHEET_COLUMN_ORDER, start=1):
            # value=None 을 ws.cell(..., value=None)으로 쓰면 기존 셀 값이 지워지지 않는다.
            # 반드시 cell.value = ... 형태로 직접 대입해야 빈 값도 덮어쓸 수 있다.
            ws.cell(row=next_row, column=col_idx).value = row.get(col_name)
        next_row += 1

    _write_notes_sheet(wb, rows)
    wb.save(target)

    notes_path = target.with_suffix(".notes.txt")
    _write_notes_txt(notes_path, rows)

    return target


def _write_notes_sheet(wb, rows: list[dict]) -> None:
    """코딩노트 시트를 생성/업데이트한다."""
    if NOTES_SHEET_NAME not in wb.sheetnames:
        ws_notes = wb.create_sheet(NOTES_SHEET_NAME)
        # 헤더
        ws_notes.cell(row=1, column=1).value = "사건 ID"
        ws_notes.cell(row=1, column=2).value = "코딩노트"
        for cell in (ws_notes.cell(row=1, column=1), ws_notes.cell(row=1, column=2)):
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
        ws_notes.column_dimensions[get_column_letter(1)].width = 18
        ws_notes.column_dimensions[get_column_letter(2)].width = 80
        start_row = 2
    else:
        ws_notes = wb[NOTES_SHEET_NAME]
        start_row = ws_notes.max_row + 1

    for row in rows:
        case_id = row.get("case_id", "")
        note = row.get("coding_note") or ""
        if not (case_id or note):
            continue
        r = start_row
        ws_notes.cell(row=r, column=1).value = case_id
        note_cell = ws_notes.cell(row=r, column=2)
        note_cell.value = note
        note_cell.alignment = Alignment(wrap_text=True, vertical="top")
        start_row += 1


def _write_notes_txt(path: Path, rows: list[dict]) -> None:
    """사건별 코딩노트를 텍스트 파일에 추가 기록한다."""
    lines = []
    for row in rows:
        case_id = row.get("case_id", "(ID없음)")
        note = row.get("coding_note") or "(코딩노트 없음)"
        lines.append(f"{'=' * 60}")
        lines.append(f"[{case_id}]")
        lines.append(note)
        lines.append("")

    if not lines:
        return

    existing = path.read_text(encoding="utf-8") if path.exists() else ""
    with path.open("a", encoding="utf-8") as f:
        if existing and not existing.endswith("\n\n"):
            f.write("\n")
        f.write("\n".join(lines) + "\n")


def notes_path_for(output_path: Path) -> Path:
    return output_path.with_suffix(".notes.txt")
